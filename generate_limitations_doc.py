from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

doc = Document()

# Title
title = doc.add_heading('InsureMEP AI Engine: Capabilities & Limitations', 0)
title.alignment = 1 # Center

def add_table_section(doc, heading, row_data):
    doc.add_heading(heading, level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    # Header
    hdr_cells = table.rows[0].cells
    headers = ['Feature / Area', 'Description', 'Business Value'] if "Capabilities" in heading else ['Weakness', 'Description', 'Real-World Impact']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.font.bold = True

    # Rows
    for col1, col2, col3 in row_data:
        row = table.add_row().cells
        row[0].text = col1
        row[1].text = col2
        row[2].text = col3

# System Capabilities Data
capabilities = [
    ("Multimodal Vision Intake", "Uses Gemini 2.0 Flash to analyze asset photos, automatically detecting asset types, visual leaks, corrosion, age, and risk metrics.", "Eliminates manual data entry for field technicians and provides instant visual risk assessment."),
    ("9-Feature RL Decision Engine", "Q-learning model trained across 16 real-world MEP trades using a comprehensive 9-feature state vector.", "Provides dynamic, trade-specific maintenance recommendations (e.g., Inspect, Repair, Defer) rather than static rules."),
    ("Warranty & Compliance Integration", "The RL engine factors in warranty coverage (reducing repair cost penalties) and overdue compliance obligations.", "Optimizes financial ROI by exploiting active warranties and avoiding compliance penalties."),
    ("Automated ETL Pipeline", "Structures the unstructured multimodal output from Gemini into a well-defined SQLite relational database schema.", "Creates a standardized data trail bridging field operations and the analytical risk engine."),
    ("Interactive Chatbot Interface", "Streamlit dashboard presenting the uploaded image, raw SQL outputs, and a live AI chatbot recommendation with Q-value breakdowns.", "Provides transparent, explainable AI decisions to facility managers.")
]

# Limitations Data
critical_gaps = [
    ("No Q-Table Persistence", "The RL agent trains from scratch for 5,000 episodes on every app boot. The Q-table is never saved to the disk.", "The \"trained\" policy changes every session. The AI has no memory and cannot accumulate learnings from actual field operations."),
    ("Synthetic Simulation", "The training environment uses randomly generated asset states and hard-coded rewards (e.g., +10/-10) instead of historical failure rates or real cost data.", "The agent learns from a disconnected, uncalibrated environment, making its real-world maintenance recommendations unreliable."),
    ("Gemini Hallucinations", "The system asks Gemini Vision to estimate non-visual metrics like risk_score, obligations_overdue, and is_under_warranty solely from a photo.", "A photo cannot reliably convey internal corrosion, compliance status, or warranty data. The confidence scores are misleading."),
    ("Zero Spatial Awareness", "When an image is uploaded, Gemini identifies the equipment type but cannot determine its physical location within the blueprint (e.g., floor, room).", "An asset photograph lacks the geospatial context required for a technician to actually locate and repair it.")
]

structural_weaknesses = [
    ("Coarse State Space", "Risk is heavily discretized into 3 simple buckets (Low, Medium, High). The 16 trades x 9 features create a massive, sparse matrix.", "Real risk is continuous. Many edge cases will never be visited during the 5,000 training episodes, leaving the Q-table severely unreliable."),
    ("Single-Asset Evaluation", "The model evaluates one asset at a time, asking \"What do I do with this one asset?\" rather than prioritizing across a whole facility.", "Real-world maintenance requires portfolio-level optimization under a fixed budget, ranking 50+ assets simultaneously."),
    ("No Real ROI / Cost Model", "Actions like urgent_repair and schedule_maintenance have different financial costs, but the reward function evaluates them using static integers.", "The RL model does not perform actual dollar-denominated Return on Investment (ROI) optimization for the facility owner."),
    ("In-Memory Database", "SQLite runs in :memory: mode, meaning all historical data is permanently wiped clean the moment the application reloads.", "The system cannot track an asset's degradation or maintenance history over time, defeating the core purpose of a CMMS.")
]

operational_issues = [
    ("No Human Feedback Loop", "There is no mechanism for the RL model to learn from a technician correcting a poor AI recommendation.", "The AI cannot improve its decision-making accuracy based on human expertise."),
    ("Hardcoded Compliance", "The compliance_flag is hardcoded to 1 in the state vector matrix.", "The compliance tracking logic is currently non-functional scaffolding."),
    ("UI Freezes on Boot", "Running 5,000 RL training episodes synchronously inside @st.cache_resource freezes the initialization of Streamlit.", "Results in a poor, unresponsive user experience during the initial application load."),
    ("No JSON Fallbacks", "If Gemini hallucinates an invalid JSON format (e.g., conversational filler), json.loads() will crash the app.", "The application is fragile and lacks robust fallback defaults for API parsing failures."),
    ("Single-Tenant Setup", "The schema lacks building_id or user_id columns, and cannot distinguish between different clients.", "The exact same system cannot reliably scale to handle multiple buildings or organizations.")
]

# Add sections
add_table_section(doc, '🌟 Current System Capabilities (Existing Features)', capabilities)
doc.add_paragraph()
doc.add_page_break()

doc.add_heading('System Limitations & Improvement Areas', level=0)
add_table_section(doc, '🔴 Critical Gaps (Requires Immediate Attention)', critical_gaps)
doc.add_paragraph()
add_table_section(doc, '🟡 Structural Weaknesses (Architectural Limitations)', structural_weaknesses)
doc.add_paragraph()
add_table_section(doc, '🟢 Operational & Engineering Issues (UX & Scalability)', operational_issues)

# Save Document
doc.save('InsureMEP_System_Assessment.docx')
print("Word document generated successfully.")
