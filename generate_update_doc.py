import sys
import subprocess
import os

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document

doc = Document()
doc.add_heading('InsureMEP AI Asset Engine - Phase 3 & 4 Update Report', 0)

doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph('We have successfully evolved the InsureMEP AI-driven maintenance engine to incorporate deep multi-modal agentic vision, strict location tracking, and critical SOP compliance mechanics. These updates directly address the onboarding and data integrity bottlenecks discussed in recent stakeholder planning sessions.')

doc.add_heading('2. Phase 3: Agentic Grid Vision & Deep Validation', level=1)

doc.add_heading('A. Agentic Grid Spatial Analysis', level=2)
p = doc.add_paragraph()
p.add_run('Concept:').bold = True
p.add_run(' Transitioned from basic image classification to a sophisticated Agentic Grid Vision model using Gemini 2.0 Flash.')
p = doc.add_paragraph()
p.add_run('Implementation:').bold = True
p.add_run(' The AI now mentally divides the uploaded asset photo into a 2x2 grid to systematically scan each quadrant. It is strictly programmed to locate Emergency Shutoff / Control Valves and determine if they are physically accessible.')

doc.add_heading('B. Emergency SOP Validation (RL Engine Integration)', level=2)
p = doc.add_paragraph()
p.add_run('State Space Expansion:').bold = True
p.add_run(' Expanded the Reinforcement Learning (Q-Learning) state vector from 9 to 10 dimensions to track "valve_accessible" status.')
p = doc.add_paragraph()
p.add_run('Penalty Mechanics:').bold = True
p.add_run(' Modeled severe SOP infractions. If a Plumbing or Fire/Life Safety asset lacks an accessible control valve, the system applies a massive (-8) reward penalty logic, instantly overriding standard maintenance suggestions and forcing an URGENT REPAIR / INSPECT ticket.')

doc.add_heading('3. Phase 4: UI/UX & Cloud System Interoperability', level=1)

doc.add_heading('A. Structural Data Enrichment (Location-First Flow)', level=2)
doc.add_paragraph('Overhauled the front-end interface (Streamlit) to require strict hierarchical location input (Building -> Floor -> Zone -> Room -> Fixture) before allowing an AI scan. This effectively eliminates the "orphan asset" problem where field technicians uploaded photos that lacked geospatial context.')

doc.add_heading('B. Explanatory AI UI (Agentic Grid Visualizer)', level=2)
doc.add_paragraph('Developed a dynamic graphic layer utilizing Python Imaging Library (PIL). When an image is analyzed, the system renders a transparent neon green 2x2 grid directly over the asset photo to visually demonstrate the AI\'s spatial processing to end users and management. Additionally, triggering the SOP violation immediately pops a giant red Blocker Modal intercepting standard flow.')

doc.add_heading('C. Production Database Sync (cloud_integration_demo.py)', level=2)
doc.add_paragraph('Developed a fully functioning Python bridge script capable of parsing real intercepted GraphQL JSON payloads right from the uci.criticalasset.tech API. The script organically downloads the live asset photos from Supabase storage into local memory, validates them through our local Agentic Vision module, and structures a viable JSON PATCH request ready to push the AI\'s compliance assessments securely back to the cloud.')

out_path = r'C:\Users\Administrator\Downloads\InsureMEP_Phase3_4_Update.docx'
os.makedirs(os.path.dirname(out_path), exist_ok=True)
doc.save(out_path)
print(f"Update report successfully saved to: {out_path}")
